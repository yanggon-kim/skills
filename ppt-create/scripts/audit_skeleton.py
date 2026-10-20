#!/usr/bin/env python3
"""Start the fact audit of a built deck: one row per checkable claim on every slide and in its notes.

Run:   <work>/venv/bin/python audit_skeleton.py DECK.pptx AUDIT.md [--source "what the claims are checked against"]

Writes a Markdown table | slide | where | claim | source location | status | with every row "todo":
  T  the title                         B  each band bullet or step caption
  F  the figure labels that carry a number (one row per slide)
  Tab each table row                   S  each headline number with its label
  C  each card, question, recap line   N  each notes sentence with a number, a %, a "×" or a reference
Fill the source location (section, figure, table, page, line) and the status for every row, merge rows that
state one fact, and keep the file next to the deck. The procedure and the statuses are in references/audit.md.
"""
import argparse
import re

from pptx import Presentation

NUM = re.compile(r"\d|%|×|§|\bFig\.|\bTable\b|\bEq\.")


ABBR = re.compile(r"(?:\b(?:Vol|Fig|Figs|Eq|Sec|No|vs|cf|al|e\.g|i\.e|approx|ca)\.|\b[A-Z]\.)$")


def sentences(text):
    """Sentences of the notes; an abbreviation ("Vol.", "Fig.", "et al.", an initial) does not end one."""
    out = []
    for piece in re.split(r"(?<=[.!?])\s+(?=[A-Z0-9\"“(])", text):
        if out and ABBR.search(out[-1]):
            out[-1] += " " + piece
        else:
            out.append(piece)
    return [s.strip() for s in out if s.strip()]


def cell(t):
    return t.replace("|", "\\|").replace("\n", " ")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck")
    ap.add_argument("out")
    ap.add_argument("--source", default="(name the source material and its version)")
    a = ap.parse_args()
    prs = Presentation(a.deck)
    rows = []
    for no, sl in enumerate(prs.slides, 1):
        figs, title = [], ""
        divider = "Section Divider" in sl.slide_layout.name  # outline, dividers, appendix index: structure only
        for sp in sl.shapes:
            nm = sp.name or ""
            if divider or nm.startswith("Slide Number"):
                continue
            if sp.has_text_frame and sp.is_placeholder and sp.placeholder_format.type is not None and "TITLE" in str(sp.placeholder_format.type):
                title = sp.text_frame.text.strip()
                rows.append((no, "T", title))
                continue
            if getattr(sp, "has_table", False) and sp.has_table:
                for r in list(sp.table.rows)[1:]:
                    rows.append((no, "Tab", " · ".join(c.text.strip() for c in r.cells)))
                continue
            if not sp.has_text_frame or not sp.text_frame.text.strip():
                continue
            text = sp.text_frame.text.strip()
            if nm in ("band", "stepcaption"):
                for p in sp.text_frame.paragraphs:
                    if p.text.strip():
                        rows.append((no, "B", p.text.strip()))
            elif nm.startswith("stat-label-"):
                rows.append((no, "S", f"{nm[len('stat-label-'):]}: {text}"))
            elif nm.startswith(("card", "q")) and nm.endswith(("-text", "-head")) or nm in ("recap", "authors", "citation"):
                rows.append((no, "C", text))
            elif nm.startswith("!!steplabel"):
                continue
            elif (":" in nm or nm.startswith("!!")) and NUM.search(text) and len(text) > 3:
                figs.append(text)
            elif sp.is_placeholder and NUM.search(text):
                rows.append((no, "C", text))
        if figs:
            rows.append((no, "F", " · ".join(dict.fromkeys(figs))))
        notes = sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else ""
        for s in sentences(notes):
            if NUM.search(s):
                rows.append((no, "N", s))
    with open(a.out, "w", encoding="utf-8") as f:
        f.write(f"# Fact audit: {a.deck.split('/')[-1]}\n\n*Checked against: {a.source}.*\n\n")
        f.write("Status: **ok** = matches the source · **fixed** = was wrong, corrected in the prototype (old → new) · "
                "**author's reading** = the author's inference or depiction, marked as such on the slide or in the notes · "
                "**background** = general knowledge, checked against (name it) · **todo** = not checked yet. "
                "Where: T title, B bullet / caption, F figure, Tab table, S stat, C card / question / recap, N notes.\n\n")
        f.write("| slide | where | claim | source location | status |\n|---|---|---|---|---|\n")
        for no, w, c in rows:
            f.write(f"| {no} | {w} | {cell(c)} |  | todo |\n")
    print(a.out, len(rows), "rows")


if __name__ == "__main__":
    main()
