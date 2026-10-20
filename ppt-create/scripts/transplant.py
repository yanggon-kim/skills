#!/usr/bin/env python3
"""Moves the slides of a raw pptxgenjs deck onto the brand template (run by deck_lib.js write()).

Run:   <work>/venv/bin/python transplant.py RAW.pptx META.json OUT.pptx [--brand DIR]
Needs: python-pptx and lxml (scripts/setup.sh puts them in <work>/venv) and the brand folder (default
       assets/ucla): Presentation-01-master.pptx, brand-01-colors.xml, brand-03-colors.xml, brand-fonts.xml.

pptxgenjs cannot open a template, so the builder writes a raw deck (10 x 5.625 in, scheme colours and theme
fonts only) plus a sidecar JSON (layout, section, Morph flag per slide; the shared geometry), and this script
transplants it:
  1. the template package is the base: its two masters, layouts, logo media and notes master come along;
     theme1 gets the Brand-01 colours, theme2 (the dark master) Brand-03, theme3 (the notes master, Office /
     Calibri in the template) Brand-01; all three the Brand fonts (Helvetica headings, Arial body); the footer
     text boxes of both masters are removed (the footer keeps only the logo and the slide number), master 1's
     title, header rule, logo and slide number move to the shared geometry (title higher, footer lower), and
     the template's docProps/app.xml (its fonts and sample-slide titles) is reset;
  2. four layouts are added inside master 1, derived from the template's own: "Header w/figure and bullets"
     (from Header w/copy, with the fixed bottom band), "Header w/storyboard" (from Header w/copy, with the
     fixed right step column), "Section Divider w/outline" (from Section Divider, title placeholder instead of
     the fixed text, rule and kicker narrowed to the left column) and "Thank You w/title" (from Thank You,
     title placeholder instead of the fixed "Thank You" text, plus a body placeholder);
  3. the template's sample slides are removed; every raw slide is re-created on its layout (title on Title
     Opener), its shapes copied, its placeholders re-pointed at the layout's (geometry and text style
     inherited), the speaker notes copied, the Morph transition (fade fallback) added to storyboard follow-on
     frames, and the sections rebuilt.
The template parts this relies on (layout file numbers, placeholder idx values, master shape names) are listed
in references/theme.md; another template needs those names changed here.
"""
import argparse
import copy
import io
import json
import os
import re
import sys
import uuid
import zipfile

from lxml import etree
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
BRAND = os.path.join(os.path.dirname(HERE), "assets", "ucla")  # replaced by --brand
TPL = os.path.join(BRAND, "Presentation-01-master.pptx")

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
A, P = "{%s}" % NS["a"], "{%s}" % NS["p"]
NSDECL = " ".join(f'xmlns:{k}="{v}"' for k, v in NS.items())
EMU = 914400

# the footer text boxes of both masters (the deck shows no footer text: only the logo and the slide number)
FOOTER_TEXT = ("Text Placeholder-left", "Text Placeholder-middle")

# ------------------------------------------------------------------ geometry shared with the JS builder
G = None  # set by main() from META.json


def e(v):
    return str(int(round(v * EMU)))


def frag(xml):
    """Parse one element written with the a:/p:/r: prefixes."""
    xml = re.sub(r"^<([\w:]+)", lambda m: f"<{m.group(1)} {NSDECL}", xml, count=1)
    return etree.fromstring(xml)


def sp_ph(name, sid, ph, box, body_pr, lst, para="<a:p><a:endParaRPr lang=\"en-US\"/></a:p>"):
    x, y, w, h = box
    return frag(
        f'<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
        f'<p:nvPr>{ph}</p:nvPr></p:nvSpPr><p:spPr><a:xfrm><a:off x="{e(x)}" y="{e(y)}"/><a:ext cx="{e(w)}" cy="{e(h)}"/>'
        f'</a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr><p:txBody>{body_pr}<a:lstStyle>{lst}</a:lstStyle>'
        f'{para}</p:txBody></p:sp>')


def rule_sp(name, sid, box, color):
    x, y, w, h = box
    return frag(
        f'<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr/><p:nvPr userDrawn="1"/></p:nvSpPr><p:spPr>'
        f'<a:xfrm><a:off x="{e(x)}" y="{e(y)}"/><a:ext cx="{e(w)}" cy="{e(h)}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/>'
        f'</a:prstGeom><a:solidFill><a:schemeClr val="{color}"/></a:solidFill><a:ln><a:noFill/></a:ln></p:spPr></p:sp>')


# ------------------------------------------------------------------ 1. the template package, rebranded
def scheme(path, tag):
    el = etree.parse(path).getroot()
    assert etree.QName(el).localname == tag, path
    return el


def set_theme(xml, colors, fonts):
    root = etree.fromstring(xml)
    te = root.find("a:themeElements", NS)
    for old, new in ((te.find("a:clrScheme", NS), colors), (te.find("a:fontScheme", NS), fonts)):
        old.getparent().replace(old, copy.deepcopy(new))
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def clean_clone(root):
    """A layout cloned from the template: drop creation ids and the designer's hidden leftovers."""
    for ext in root.iter(A + "extLst"):
        if ext.getparent().tag == P + "cNvPr" or ext.getparent().tag == A + "cNvPr":
            ext.getparent().remove(ext)
    top = root.find("p:extLst", NS)
    if top is not None:
        root.remove(top)
    tree = root.find("p:cSld/p:spTree", NS)
    for el in list(tree):
        c = el.find(".//p:cNvPr", NS)
        if c is not None and c.get("hidden") == "1":
            tree.remove(el)
    return tree


def ph_of(el):
    return el.find("./*/p:nvPr/p:ph", NS)


def drop(tree, pred):
    for el in list(tree):
        if pred(el):
            tree.remove(el)


def name_of(el):
    c = el.find("./*/p:cNvPr", NS)
    return c.get("name") if c is not None else ""


def set_xfrm(el, box):
    """Place a shape at box (x, y, w, h in inches); the shape must carry its own a:xfrm."""
    xf = el.find("p:spPr/a:xfrm", NS)
    xf.find("a:off", NS).set("x", e(box[0]))
    xf.find("a:off", NS).set("y", e(box[1]))
    xf.find("a:ext", NS).set("cx", e(box[2]))
    xf.find("a:ext", NS).set("cy", e(box[3]))


def restyle_master(xml, main):
    """Remove the footer text boxes; on the main master also move the title, header rule, logo and slide number."""
    root = etree.fromstring(xml)
    tree = root.find("p:cSld/p:spTree", NS)
    drop(tree, lambda el: name_of(el) in FOOTER_TEXT)
    assert not any(name_of(el) in FOOTER_TEXT for el in tree)
    if main:
        byname = {name_of(el): el for el in tree}
        set_xfrm(byname["Title Placeholder"], G["title"])
        set_xfrm(byname["Header rule"], G["hrule"])
        set_xfrm(byname["Brand logo"], G["footLogo"])
        num = byname["Slide Number Placeholder"]
        set_xfrm(num, G["sldNum"])
        bp = num.find("p:txBody/a:bodyPr", NS)
        bp.set("bIns", "0")
        for ch in list(bp):
            bp.remove(ch)  # no autofit: the box keeps the fixed footer height
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def layouts_xml(files):
    """The four derived layouts, as {layout name: xml bytes}."""
    out = {}
    TB = '<a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" rtlCol="0" anchor="t" anchorCtr="0"><a:noAutofit/></a:bodyPr>'

    def title_lst(sz):
        return f'<a:lvl1pPr><a:defRPr sz="{sz}"/></a:lvl1pPr>'

    # --- content: Header w/copy + fixed bottom band (bullets and 14 pt from the master body style)
    for lname, kind in (("Header w/figure and bullets", "CONTENT"), ("Header w/storyboard", "FRAME")):
        root = etree.fromstring(files["ppt/slideLayouts/slideLayout3.xml"])
        tree = clean_clone(root)
        root.find("p:cSld", NS).set("name", lname)
        drop(tree, lambda el: ph_of(el) is not None and (ph_of(el).get("idx") in ("13", "15") or ph_of(el).get("type") == "dt"))
        for el in tree:
            ph = ph_of(el)
            if ph is not None and ph.get("type") == "title":
                set_xfrm(el, G["title"])
                el.find("p:txBody/a:lstStyle", NS).append(frag(title_lst(G["titlePt"] * 100)))
        if kind == "CONTENT":
            tree.append(rule_sp("band rule", 40, G["rule"], "accent4"))
            lst = (f'<a:lvl1pPr><a:spcBef><a:spcPts val="{G["bandSpcBef"] * 100}"/></a:spcBef>'
                   f'<a:defRPr sz="{int(G["textPt"] * 100)}"/></a:lvl1pPr>')
            tree.append(sp_ph("band", 41, '<p:ph type="body" sz="quarter" idx="13"/>', G["band"], TB, lst))
        else:
            lst = ('<a:lvl1pPr marL="0" indent="0"><a:spcBef><a:spcPts val="0"/></a:spcBef><a:buNone/>'
                   f'<a:defRPr sz="{int(G["textPt"] * 100)}"/></a:lvl1pPr>')
            tree.append(sp_ph("stepcaption", 41, '<p:ph type="body" sz="quarter" idx="13"/>', G["stepcaption"], TB, lst))
        out[lname] = root

    # --- section divider: title placeholder instead of the fixed text; rule and kicker in the left column
    root = etree.fromstring(files["ppt/slideLayouts/slideLayout2.xml"])
    tree = clean_clone(root)
    root.find("p:cSld", NS).set("name", "Section Divider w/outline")
    drop(tree, lambda el: name_of(el).startswith("TextBox") or (ph_of(el) is not None and ph_of(el).get("type") == "dt"))
    lw = G["dividerColW"]
    for el in tree:
        if name_of(el) == "Header rule" or (ph_of(el) is not None and ph_of(el).get("idx") == "20"):
            el.find("p:spPr/a:xfrm/a:ext", NS).set("cx", e(lw))
    tree.append(sp_ph("Title", 42, '<p:ph type="title"/>', (0.70, 2.05, lw, 0.61),
                      '<a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" rtlCol="0" anchor="b" anchorCtr="0"><a:noAutofit/></a:bodyPr>',
                      '<a:lvl1pPr><a:defRPr sz="3600" b="1"/></a:lvl1pPr>'))
    out["Section Divider w/outline"] = root

    # --- closing: Thank You with a title placeholder instead of the fixed text, and a body placeholder
    root = etree.fromstring(files["ppt/slideLayouts/slideLayout16.xml"])
    tree = clean_clone(root)
    root.find("p:cSld", NS).set("name", "Thank You w/title")
    drop(tree, lambda el: name_of(el).startswith("TextBox") or (ph_of(el) is not None and ph_of(el).get("type") == "dt"))
    tree.append(sp_ph("Title", 42, '<p:ph type="title"/>', (0.70, 2.05, 8.50, 0.61),
                      '<a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" rtlCol="0" anchor="b" anchorCtr="0"><a:noAutofit/></a:bodyPr>',
                      '<a:lvl1pPr><a:defRPr sz="3600" b="1"/></a:lvl1pPr>'))
    tree.append(sp_ph("recap", 43, '<p:ph type="body" sz="quarter" idx="20"/>', G["closingBody"], TB,
                      f'<a:lvl1pPr marL="0" indent="0"><a:spcBef><a:spcPts val="0"/></a:spcBef><a:buNone/><a:defRPr sz="{int(G["textPt"] * 100)}"/></a:lvl1pPr>'))
    out["Thank You w/title"] = root
    return {k: etree.tostring(v, xml_declaration=True, encoding="UTF-8", standalone=True) for k, v in out.items()}


def branded_template(n_slides):
    zin = zipfile.ZipFile(TPL)
    files = {n: zin.read(n) for n in zin.namelist()}
    c01 = scheme(os.path.join(BRAND, "brand-01-colors.xml"), "clrScheme")
    c03 = scheme(os.path.join(BRAND, "brand-03-colors.xml"), "clrScheme")
    fonts = scheme(os.path.join(BRAND, "brand-fonts.xml"), "fontScheme")
    files["ppt/theme/theme1.xml"] = set_theme(files["ppt/theme/theme1.xml"], c01, fonts)
    files["ppt/theme/theme2.xml"] = set_theme(files["ppt/theme/theme2.xml"], c03, fonts)  # the dark master
    files["ppt/theme/theme3.xml"] = set_theme(files["ppt/theme/theme3.xml"], c01, fonts)  # the notes master (was Office/Calibri)
    # the template's app.xml lists its own fonts and sample-slide titles; PowerPoint rewrites this part on save
    files["docProps/app.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Properties xmlns="http://schemas.openxmlformats.org/'
        'officeDocument/2006/extended-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        '<TotalTime>0</TotalTime><Application>Microsoft Office PowerPoint</Application><PresentationFormat>On-screen Show (16:9)'
        f'</PresentationFormat><Slides>{n_slides}</Slides><Notes>{n_slides}</Notes><HiddenSlides>0</HiddenSlides><ScaleCrop>false'
        '</ScaleCrop><LinksUpToDate>false</LinksUpToDate><SharedDoc>false</SharedDoc><HyperlinksChanged>false</HyperlinksChanged>'
        '<AppVersion>16.0000</AppVersion></Properties>').encode()
    for m in ("ppt/slideMasters/slideMaster1.xml", "ppt/slideMasters/slideMaster2.xml"):
        files[m] = restyle_master(files[m], m.endswith("1.xml"))

    # register the derived layouts inside master 1
    new = layouts_xml(files)
    n0 = max(int(m.group(1)) for f in files for m in [re.match(r"ppt/slideLayouts/slideLayout(\d+)\.xml$", f)] if m)
    mrels = etree.fromstring(files["ppt/slideMasters/_rels/slideMaster1.xml.rels"])
    master = etree.fromstring(files["ppt/slideMasters/slideMaster1.xml"])
    ctypes = etree.fromstring(files["[Content_Types].xml"])
    ids = [int(v) for f in ("ppt/presentation.xml", "ppt/slideMasters/slideMaster1.xml", "ppt/slideMasters/slideMaster2.xml")
           for v in re.findall(rb'<p:sld(?:Master|Layout)Id id="(\d+)"', files[f])]
    nid = max(ids) + 1
    rid = max(int(r.get("Id")[3:]) for r in mrels) + 1
    lst = master.find("p:sldLayoutIdLst", NS)
    for i, (lname, xml) in enumerate(new.items()):
        part = f"slideLayout{n0 + 1 + i}.xml"
        files[f"ppt/slideLayouts/{part}"] = xml
        files[f"ppt/slideLayouts/_rels/{part}.rels"] = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/'
            'package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
            'relationships/slideMaster" Target="../slideMasters/slideMaster1.xml"/></Relationships>').encode()
        etree.SubElement(mrels, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                         Id=f"rId{rid}", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout",
                         Target=f"../slideLayouts/{part}")
        el = etree.SubElement(lst, P + "sldLayoutId")
        el.set("id", str(nid))
        el.set("{%s}id" % NS["r"], f"rId{rid}")
        etree.SubElement(ctypes, "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                         PartName=f"/ppt/slideLayouts/{part}",
                         ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml")
        nid += 1
        rid += 1
    for k, el in (("ppt/slideMasters/_rels/slideMaster1.xml.rels", mrels), ("ppt/slideMasters/slideMaster1.xml", master),
                  ("[Content_Types].xml", ctypes)):
        files[k] = etree.tostring(el, xml_declaration=True, encoding="UTF-8", standalone=True)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for n in zin.namelist() + [n for n in files if n not in zin.namelist()]:
            z.writestr(n, files[n])
    return buf, files["ppt/media/image1.png"]


# ------------------------------------------------------------------ 2. slides
LAYOUT = {"TITLE": "Title Opener", "DIVIDER": "Section Divider w/outline", "CONTENT": "Header w/figure and bullets",
          "FRAME": "Header w/storyboard", "CLOSING": "Thank You w/title"}
# raw placeholder (raw layout, type) -> the target layout's placeholder
PHMAP = {("TITLE", "title"): ("body", "31"), ("DIVIDER", "title"): ("title", None), ("DIVIDER", "body"): ("body", "20"),
         ("CONTENT", "title"): ("title", None), ("CONTENT", "body"): ("body", "13"),
         ("FRAME", "title"): ("title", None), ("FRAME", "body"): ("body", "13"),
         ("CLOSING", "title"): ("title", None), ("CLOSING", "body"): ("body", "20")}
# text style of these placeholders comes from the layout (pptxgenjs writes size and colour on every run)
INHERIT = {("CONTENT", "title"), ("FRAME", "title"), ("CONTENT", "body"), ("FRAME", "body"), ("DIVIDER", "title"),
           ("CLOSING", "title"), ("CLOSING", "body")}
MORPH = ('<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
         '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">'
         '<p:transition spd="slow"><p159:morph option="byObject"/></p:transition></mc:Choice>'
         '<mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')
SLDNUM = ('<p:sp><p:nvSpPr><p:cNvPr id="{id}" name="Slide Number Placeholder"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
          '<p:nvPr><p:ph type="sldNum" sz="quarter" idx="19"/></p:nvPr></p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/>'
          '<a:p><a:fld id="{{B6238B5B-F19C-E947-A0BC-87BD7983F871}}" type="slidenum"><a:rPr lang="en-US"/><a:t>{n}</a:t></a:fld>'
          '<a:endParaRPr lang="en-US"/></a:p></p:txBody></p:sp>')


def strip_text_style(sp, keep_para=False):
    body = sp.find("p:txBody", NS)
    bp = body.find("a:bodyPr", NS)
    for k in list(bp.attrib):
        del bp.attrib[k]
    for ch in list(bp):
        bp.remove(ch)
    for p in body.findall("a:p", NS):
        if not keep_para:
            for ppr in p.findall("a:pPr", NS):
                p.remove(ppr)
        for r in p.iter(A + "rPr", A + "endParaRPr"):
            for k in ("sz", "b"):
                r.attrib.pop(k, None)
            for f in r.findall("a:solidFill", NS):
                r.remove(f)


def as_placeholder(sp, ptype, idx):
    nv = sp.find("p:nvSpPr", NS)
    nvpr = nv.find("p:nvPr", NS)
    for old in nvpr.findall("p:ph", NS):
        nvpr.remove(old)
    ph = etree.Element(P + "ph", type=ptype)
    if idx is not None:
        ph.set("idx", idx)
    nvpr.insert(0, ph)
    c = nv.find("p:cNvSpPr", NS)
    c.attrib.pop("txBox", None)
    if c.find("a:spLocks", NS) is None:
        etree.SubElement(c, A + "spLocks", noGrp="1")


def main():
    global G, BRAND, TPL
    ap = argparse.ArgumentParser(description="Transplant a raw deck onto the brand template.")
    ap.add_argument("raw")
    ap.add_argument("meta")
    ap.add_argument("out")
    ap.add_argument("--brand", default=BRAND, help="folder with the template and theme files")
    a = ap.parse_args()
    raw_path, meta_path, out_path = a.raw, a.meta, a.out
    BRAND = os.path.abspath(a.brand)
    TPL = os.path.join(BRAND, "Presentation-01-master.pptx")
    meta = json.load(open(meta_path, encoding="utf-8"))
    G = meta["geometry"]
    tpl_buf, logo_png = branded_template(len(meta["slides"]))
    prs = Presentation(tpl_buf)
    raw = Presentation(raw_path)
    assert len(raw.slides) == len(meta["slides"]), (len(raw.slides), len(meta["slides"]))

    # remove the template's sample slides (their parts are dropped on save: nothing references them)
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        prs.part.drop_rel(sid.rId)
        lst.remove(sid)
    layouts = {l.name: l for l in prs.slide_masters[0].slide_layouts}

    sections, new_ids = [], []
    for no, (rs, m) in enumerate(zip(raw.slides, meta["slides"]), start=1):
        kind = m["layout"]
        slide = prs.slides.add_slide(layouts[LAYOUT[kind]])
        tree = slide.shapes._spTree
        keep_logo = None
        for el in list(tree)[2:]:
            ph = ph_of(el)
            if kind == "TITLE" and ph is not None and ph.get("idx") == "24":
                keep_logo = el
                continue
            tree.remove(el)
        if keep_logo is not None:
            phobj = [s for s in slide.placeholders if s.placeholder_format.idx == 24][0]
            pic = phobj.insert_picture(io.BytesIO(logo_png))
            pic.crop_left = pic.crop_right = pic.crop_top = pic.crop_bottom = 0
            lg = meta["geometry"]["logo"]
            pic.left, pic.top, pic.width, pic.height = (int(round(v * EMU)) for v in lg)
            pic.name = "UCLA logo"
        rtree = rs.shapes._spTree
        for el in list(rtree)[2:]:
            ph = ph_of(el)
            if ph is not None and ph.get("type") == "sldNum":
                continue
            el = copy.deepcopy(el)
            nm = name_of(el)
            for blip in el.iter(A + "blip"):  # a picture: carry its image part over to the new slide
                rid = blip.get("{%s}embed" % NS["r"])
                if rid:
                    blob = rs.part.related_part(rid).blob
                    _, new_rid = slide.part.get_or_add_image_part(io.BytesIO(blob))
                    blip.set("{%s}embed" % NS["r"], new_rid)
            if ph is not None:
                key = (kind, ph.get("type") or "body")
                ptype, idx = PHMAP[key]
                as_placeholder(el, ptype, idx)
                sppr = el.find("p:spPr", NS)
                for ch in list(sppr):
                    sppr.remove(ch)
                if key in INHERIT:
                    strip_text_style(el)
                if key == ("TITLE", "title"):  # the subtitle paragraph is not bold (pptxgenjs omits b="0")
                    for p_ in el.findall("p:txBody/a:p", NS)[1:]:
                        for r_ in p_.iter(A + "rPr"):
                            r_.set("b", "0")
                if key == ("CONTENT", "body"):
                    el.find("p:nvSpPr/p:cNvPr", NS).set("name", "band")
                    for r_ in el.iter(A + "r"):  # the appendix pointer keeps its emphasis: bold, accent6
                        if (r_.findtext("a:t", namespaces=NS) or "").startswith(" → appendix "):
                            rpr = r_.find("a:rPr", NS)
                            rpr.set("b", "1")
                            fill = etree.SubElement(rpr, A + "solidFill")
                            etree.SubElement(fill, A + "schemeClr", val="accent6")
                            rpr.insert(0, fill)
                elif key == ("FRAME", "body"):
                    el.find("p:nvSpPr/p:cNvPr", NS).set("name", "stepcaption")
            elif nm.startswith("ph:"):  # a text box the JS asks to become a template placeholder (keeps its box)
                _, ptype, idx = nm.split(":")
                as_placeholder(el, ptype, idx)
            tree.append(el)
        if m["slideNumber"]:
            tree.append(frag(SLDNUM.format(id=0, n=no)))
        # unique shape ids
        for i, c in enumerate(tree.iter(P + "cNvPr"), start=1):
            c.set("id", str(i))
        if m["morph"]:
            sld = slide._element
            cm = sld.find("p:clrMapOvr", NS)
            assert cm is not None and sld.find("p:transition", NS) is None
            cm.addnext(frag(MORPH))
        notes = rs.notes_slide.notes_text_frame.text if rs.has_notes_slide else ""
        slide.notes_slide.notes_text_frame.text = notes
        new_ids.append(slide.slide_id)
        if not sections or sections[-1][0] != m["section"]:
            sections.append((m["section"], []))
        sections[-1][1].append(slide.slide_id)

    # sections
    pres = prs.part._element
    ext = pres.find("p:extLst", NS)
    if ext is None:
        ext = etree.SubElement(pres, P + "extLst")
    sec = ('<p:ext uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"><p14:sectionLst xmlns:p14="http://schemas.microsoft.com/office/'
           'powerpoint/2010/main">' + "".join(
               f'<p14:section name="{n}" id="{{{str(uuid.uuid5(uuid.NAMESPACE_URL, "ppt-create/" + meta["title"] + "/" + n)).upper()}}}"><p14:sldIdLst>'
               + "".join(f'<p14:sldId id="{i}"/>' for i in ids) + "</p14:sldIdLst></p14:section>" for n, ids in sections)
           + "</p14:sectionLst></p:ext>")
    ext.insert(0, frag(sec))

    cp = prs.core_properties
    cp.title, cp.subject, cp.author = meta["title"], meta["subject"], meta["author"]
    cp.last_modified_by = meta["author"]
    cp.revision = 1
    prs.save(out_path)
    print(f"{out_path}: {len(new_ids)} slides on the brand template; layouts used:",
          ", ".join(sorted({LAYOUT[m['layout']] for m in meta['slides']})))


if __name__ == "__main__":
    main()
