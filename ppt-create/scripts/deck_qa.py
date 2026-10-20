#!/usr/bin/env python3
"""Structural and fit checks for a deck built by deck_lib.js (run after the build).

Run:   <work>/venv/bin/python deck_qa.py DECK.pptx [<work>/DECK.meta.json]
Needs: python-pptx, lxml and Pillow (scripts/setup.sh) and the Liberation Sans fonts (metric stand-ins for Arial
       and Helvetica; the renders use them too). Without the sidecar JSON only the slide-level checks run.

Checks, each printed as PASS / FAIL with the offending slides; exit status 1 on any FAIL:
  footer    no footer text on any master or layout; the logo and slide number at the shared geometry
  geometry  title, band, step-column and caption placeholders at one geometry on every layout that has them
  floors    no slide text below its floor (12 pt; bullets / step captions 16 pt; tables 13 pt, a dense table 12)
  tags      no provenance text on a slide ("Source:", "Redrawn", "Adapted from", "schematic", ...); a restricted
            phrase (sidecar qa.restricted) only inside the named figures
  morph     one Morph transition per storyboard follow-on frame; the !!-names match frame to frame
  notes     every slide has speaker notes
  structure title first, outline second (it opens the Summary part), the summary slide third with a title
            starting "Summary", a divider opening every other part, one closing slide, appendix slides only
            after the closing slide
  fit       every wrapping text box holds its text (greedy word wrap with Liberation Sans at 97 % of the width),
            every title is one line, every table ends above the band
  overlap   no two texts overlap (each text's ink: its measured width placed by alignment, its lines placed by
            the anchor, trimmed to the glyphs, so a box's empty padding never counts)
Warnings (printed as WARN, exit status unchanged; look at the render and fix or accept each):
  lines     a line or connector passes through a text's ink (gridlines, arrows, curves over labels)
  edges     a text sits mostly inside a drawn shape but runs out over its edge (a label wider than its box)
  whitespace  the share of the content area (frame area on storyboard slides) farther than 0.25 in from
            every element is above 25 % for a table slide, 42 % for a figure slide or a frame
The table-height estimate runs about 0.1 in high: when a table fails by that margin, trust the render.
"""
import json
import os
import re
import subprocess
import sys

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Emu

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "p159": "http://schemas.microsoft.com/office/powerpoint/2015/09/main"}
EMU = 914400


def _font_file(bold):
    """Liberation Sans (Arial / Helvetica metrics): the usual paths, else fontconfig."""
    name = "LiberationSans-Bold.ttf" if bold else "LiberationSans-Regular.ttf"
    for d in ("/usr/share/fonts/truetype/liberation2", "/usr/share/fonts/truetype/liberation",
              "/usr/share/fonts/liberation-sans", "/usr/share/fonts/liberation"):
        if os.path.isfile(os.path.join(d, name)):
            return os.path.join(d, name)
    try:
        f = subprocess.run(["fc-match", "-f", "%{file}", "Liberation Sans:bold" if bold else "Liberation Sans"],
                           capture_output=True, text=True).stdout.strip()
        if f:
            return f
    except OSError:
        pass
    sys.exit("Liberation Sans not found: install fonts-liberation (the fit check measures text with it)")


FONT = {False: _font_file(False), True: _font_file(True)}
_fonts = {}
TAGS = r"Paper's|After the paper|[Ss]chematic|Redrawn|Adapted from|[Ss]ource:|Courtesy of|Reproduced from"
FOOTER_WORDS = ("edit in Slide Master",)
WS_REACH = 0.25  # in: an element "uses" the area within this distance of it (gaps inside a figure count as used)
# warn when more than this share of the area is empty. Calibrated on the approved reference deck
# (assets/reference/p3llm-intro.pptx): its frames reach 40.2 %, its tables 3 %, its figure slides 39.5 %
# except slide 35 (49.5 %, a tall dot plot with bands left and right), which is its only warning.
WS_WARN = {"table": 0.25, "content": 0.42, "frame": 0.42}
OVERLAP_MIN = 0.02  # in: two text inks that overlap by more than this in both directions FAIL


def font(sz, bold):
    k = (round(sz * 4), bold)
    if k not in _fonts:
        _fonts[k] = ImageFont.truetype(FONT[bold], max(1, round(sz * 4)))  # 4x for sub-point accuracy
    return _fonts[k]


def width_in(text, sz, bold):
    return font(sz, bold).getlength(text) / 4 / 72


def wrap_lines(runs, width):
    """runs: [(text, sz, bold)] of one paragraph -> number of lines under greedy word wrap."""
    words = []
    for t, sz, b in runs:
        for i, w in enumerate(re.split(r"(\s+)", t)):
            if w:
                words.append((w, sz, b))
    lines, cur = 1, 0.0
    for w, sz, b in words:
        wl = width_in(w, sz, b)
        if w.isspace():
            cur += wl
            continue
        if cur > 0 and cur + wl > width:
            lines += 1
            cur = wl
        else:
            cur += wl
    return lines


def emu_in(v):
    return int(v) / EMU


class Inherit:
    """Placeholder geometry and text defaults from the layout, then the master."""

    def __init__(self, slide):
        self.layout = slide.slide_layout
        self.master = self.layout.slide_master

    def ph(self, ph_el):
        typ, idx = ph_el.get("type") or "body", ph_el.get("idx")
        for part in (self.layout, self.master):
            for sp in part.placeholders:
                p = sp._element.find(".//p:nvPr/p:ph", NS)
                if (idx is not None and p.get("idx") == idx) or (idx is None and (p.get("type") or "body") == typ):
                    yield sp._element
            # the master's title / body placeholders match by type
            if part is self.master:
                for sp in part.placeholders:
                    p = sp._element.find(".//p:nvPr/p:ph", NS)
                    if (p.get("type") or "body") == typ:
                        yield sp._element


def first(it, f):
    for el in it:
        v = f(el)
        if v is not None:
            return v
    return None


def box_of(el, inh):
    xf = el.find("p:spPr/a:xfrm", NS)
    ph = el.find(".//p:nvPr/p:ph", NS)
    if xf is None and ph is not None:
        xf = first(inh.ph(ph), lambda s: s.find("p:spPr/a:xfrm", NS))
    if xf is None:
        return None
    o, x = xf.find("a:off", NS), xf.find("a:ext", NS)
    return [emu_in(o.get("x")), emu_in(o.get("y")), emu_in(x.get("cx")), emu_in(x.get("cy"))]


def body_attr(el, inh, k, default):
    v = el.find("p:txBody/a:bodyPr", NS).get(k)
    ph = el.find(".//p:nvPr/p:ph", NS)
    if v is None and ph is not None:
        v = first(inh.ph(ph), lambda s: s.find("p:txBody/a:bodyPr", NS).get(k) if s.find("p:txBody/a:bodyPr", NS) is not None else None)
    return default if v is None else v


def lst_val(el, inh, path, attr):
    """A lvl1pPr value: the shape's lstStyle, then the layout / master placeholder, then the master text style."""
    srcs = [el]
    ph = el.find(".//p:nvPr/p:ph", NS)
    if ph is not None:
        srcs += list(inh.ph(ph))
    for s in srcs:
        n = s.find(f"p:txBody/a:lstStyle/a:lvl1pPr{path}", NS)
        if n is not None and n.get(attr) is not None:
            return n.get(attr)
    if ph is not None:
        style = "p:titleStyle" if (ph.get("type") == "title") else "p:bodyStyle"
        n = inh.master._element.find(f"p:txStyles/{style}/a:lvl1pPr{path}", NS)
        if n is not None:
            return n.get(attr)
    return None


def run_size(r, el, inh, ptype):
    rpr = r.find("a:rPr", NS)
    if rpr is not None and rpr.get("sz"):
        return int(rpr.get("sz")) / 100
    v = lst_val(el, inh, "/a:defRPr", "sz")
    return int(v) / 100 if v else 18.0


def run_bold(r, el, inh):
    rpr = r.find("a:rPr", NS)
    if rpr is not None and rpr.get("b") is not None:
        return rpr.get("b") == "1"
    v = lst_val(el, inh, "/a:defRPr", "b")
    return v == "1"


def seg_hits_rect(x1, y1, x2, y2, r):
    """Liang-Barsky: does the segment (x1, y1)-(x2, y2) pass through the rectangle r = [x0, y0, x1, y1]?"""
    dx, dy = x2 - x1, y2 - y1
    t0, t1 = 0.0, 1.0
    for p, q in ((-dx, x1 - r[0]), (dx, r[2] - x1), (-dy, y1 - r[1]), (dy, r[3] - y1)):
        if p == 0:
            if q < 0:
                return False
            continue
        t = q / p
        if p < 0:
            t0 = max(t0, t)
        else:
            t1 = min(t1, t)
        if t0 > t1:
            return False
    return True


def seg_dist(px, py, x1, y1, x2, y2):
    dx, dy = x2 - x1, y2 - y1
    L = dx * dx + dy * dy
    t = 0.0 if L == 0 else max(0.0, min(1.0, ((px - x1) * dx + (py - y1) * dy) / L))
    return ((px - x1 - t * dx) ** 2 + (py - y1 - t * dy) ** 2) ** 0.5


def rect_dist(px, py, r):
    ddx = max(r[0] - px, 0.0, px - r[2])
    ddy = max(r[1] - py, 0.0, py - r[3])
    return (ddx * ddx + ddy * ddy) ** 0.5


def empty_fraction(area, blocks, segs, reach=WS_REACH, cell=0.05):
    """Share of the content area farther than `reach` inches from every element (shape box, text ink, line).
    Gaps inside a figure (between boxes, around labels) count as used; bands and corners left bare do not."""
    ax, ay, aw, ah = area
    nx, ny = int(round(aw / cell)), int(round(ah / cell))
    used = bytearray(nx * ny)

    def mark(x0, y0, x1, y1, inside):
        i0, i1 = max(0, int((x0 - reach - ax) / cell)), min(nx - 1, int((x1 + reach - ax) / cell))
        j0, j1 = max(0, int((y0 - reach - ay) / cell)), min(ny - 1, int((y1 + reach - ay) / cell))
        for j in range(j0, j1 + 1):
            cy = ay + (j + 0.5) * cell
            for i in range(i0, i1 + 1):
                k = j * nx + i
                if not used[k] and inside(ax + (i + 0.5) * cell, cy):
                    used[k] = 1

    for r in blocks:
        mark(r[0], r[1], r[2], r[3], lambda px, py, r=r: rect_dist(px, py, r) <= reach)
    for s in segs:
        mark(min(s[0], s[2]), min(s[1], s[3]), max(s[0], s[2]), max(s[1], s[3]),
             lambda px, py, s=s: seg_dist(px, py, *s) <= reach)
    return 1 - sum(used) / (nx * ny)


def main():
    path = sys.argv[1]
    meta = json.load(open(sys.argv[2], encoding="utf-8")) if len(sys.argv) > 2 else None
    G = meta["geometry"] if meta else None
    qa = meta.get("qa", {}) if meta else {}
    tags = re.compile("|".join([TAGS] + list(qa.get("extraTags", []))))
    footer_words = FOOTER_WORDS + tuple(w for w in qa.get("footerWords", []) if w)
    restricted = qa.get("restricted", [])
    table_floors = qa.get("tableFloors", {})
    prs = Presentation(path)
    res = {k: [] for k in ("footer", "geometry", "floors", "tags", "morph", "notes", "structure", "fit", "overlap")}
    warns = {k: [] for k in ("lines", "edges", "whitespace")}
    fit_notes, ws_notes = [], []
    areas = {"Header w/figure and bullets": G["area"], "Header w/storyboard": G["farea"]} if G else {}

    # footer text and footer geometry
    for m in prs.slide_masters:
        for part in [m] + list(m.slide_layouts):
            txt = " ".join(t.text or "" for t in part._element.iter("{%s}t" % NS["a"]))
            for w in footer_words:
                if len(w) >= 4 and w in txt:
                    res["footer"].append(f"{part.name if hasattr(part, 'name') else 'master'}: {w}")
    m1 = prs.slide_masters[0]._element
    if G:
        for nm, key in (("Brand logo", "footLogo"), ("Slide Number Placeholder", "sldNum"), ("Header rule", "hrule"), ("Title Placeholder", "title")):
            el = [c for c in m1.iter("{%s}cNvPr" % NS["p"]) if c.get("name") == nm][0].getparent().getparent()
            xf = el.find("p:spPr/a:xfrm", NS)
            got = [emu_in(xf.find("a:off", NS).get("x")), emu_in(xf.find("a:off", NS).get("y")),
                   emu_in(xf.find("a:ext", NS).get("cx")), emu_in(xf.find("a:ext", NS).get("cy"))]
            if any(abs(a - b) > 0.005 for a, b in zip(got, G[key])):
                res["geometry"].append(f"master {nm} at {[round(v, 3) for v in got]} != {G[key]}")

    # derived layouts: one geometry each
    if G:
        for lay in prs.slide_masters[0].slide_layouts:
            for sp in lay.placeholders:
                p = sp._element.find(".//p:nvPr/p:ph", NS)
                want = None
                if lay.name in ("Header w/figure and bullets", "Header w/storyboard") and p.get("type") == "title":
                    want = G["title"]
                elif lay.name == "Header w/figure and bullets" and p.get("idx") == "13":
                    want = G["band"]
                elif lay.name == "Header w/storyboard" and p.get("idx") == "13":
                    want = G["stepcaption"]
                if want:
                    got = [Emu(v).inches for v in (sp.left, sp.top, sp.width, sp.height)]
                    if any(abs(a - b) > 0.005 for a, b in zip(got, want)):
                        res["geometry"].append(f"{lay.name} {p.get('type')}/{p.get('idx')} at {got} != {want}")

    morph_n = 0
    story_names = {}
    for no, slide in enumerate(prs.slides, start=1):
        inh = Inherit(slide)
        lay = slide.slide_layout.name
        sld = slide._element
        if sld.find(".//p159:morph", NS) is not None:
            morph_n += 1
        names = [c.get("name") for c in sld.iter("{%s}cNvPr" % NS["p"])]
        if lay == "Header w/storyboard":
            story_names[no] = (slide.shapes.title.text, sorted(n for n in names if n.startswith("!!")))
        for c in sld.iter("{%s}cNvPr" % NS["p"]):
            if c.get("name") in ("srctag", "!!srctag"):
                res["tags"].append(f"s{no}: shape {c.get('name')}")
        # step column: dots and step label at one place
        texts, segs, blocks = [], [], []  # text inks, line segments, other content boxes (inches, x0 y0 x1 y1)
        has_table = False
        boxes = []  # drawn shapes (rects, ovals), for the label-over-an-edge warning
        for sp in slide.shapes:
            el = sp._element
            nm = sp.name
            is_ph = el.find(".//p:nvPr/p:ph", NS) is not None
            in_col = nm.startswith(("!!dot", "!!steplabel"))  # the step column is fixed text, not content
            pgeom = el.find("p:spPr/a:prstGeom", NS)
            prst = pgeom.get("prst") if pgeom is not None else None
            if prst == "line" or el.tag.endswith("}cxnSp"):
                xf = el.find("p:spPr/a:xfrm", NS)
                if xf is not None:
                    x, y = emu_in(xf.find("a:off", NS).get("x")), emu_in(xf.find("a:off", NS).get("y"))
                    w, h = emu_in(xf.find("a:ext", NS).get("cx")), emu_in(xf.find("a:ext", NS).get("cy"))
                    x1, x2 = (x + w, x) if xf.get("flipH") == "1" else (x, x + w)
                    y1, y2 = (y + h, y) if xf.get("flipV") == "1" else (y, y + h)
                    segs.append((x1, y1, x2, y2, nm))
                continue
            if not is_ph and not in_col and not (hasattr(sp, "has_table") and sp.has_table):
                spr = el.find("p:spPr", NS)
                ln = spr.find("a:ln", NS) if spr is not None else None
                # drawn: a picture, a fill, or an outline with its own colour (a bare <a:ln/> on a text box draws nothing)
                visible = el.tag.endswith("}pic") or (spr is not None and (spr.find("a:solidFill", NS) is not None or
                                                                             (ln is not None and ln.find("a:solidFill", NS) is not None)))
                bx = box_of(el, inh) if spr is not None else None
                if visible and bx:
                    blocks.append([bx[0], bx[1], bx[0] + bx[2], bx[1] + bx[3]])
                    if not el.tag.endswith("}pic"):
                        boxes.append((blocks[-1], nm))
            if G and nm == "!!steplabel" and abs(Emu(sp.top).inches - G.get("stepLabelY", 1.14)) > 0.005:
                res["geometry"].append(f"s{no}: steplabel y {Emu(sp.top).inches}")
            # tables
            if sp.has_table if hasattr(sp, "has_table") else False:
                floor = table_floors.get(nm, 13)
                tot = 0.0
                for row in sp.table.rows:
                    hmax = 0.0
                    for cell in row.cells:
                        tc = cell._tc
                        pr = tc.find("a:tcPr", NS)
                        ml = emu_in(pr.get("marL", 91440)) if pr is not None else 0.1
                        mr = emu_in(pr.get("marR", 91440)) if pr is not None else 0.1
                        mt = emu_in(pr.get("marT", 45720)) if pr is not None else 0.05
                        mb = emu_in(pr.get("marB", 45720)) if pr is not None else 0.05
                        w = Emu(sp.table.columns[list(row.cells).index(cell)].width).inches - ml - mr
                        hh = 0.0
                        for p in tc.findall("a:txBody/a:p", NS):
                            runs = []
                            for r in p.findall("a:r", NS):
                                rp = r.find("a:rPr", NS)
                                sz = int(rp.get("sz")) / 100 if rp is not None and rp.get("sz") else 18
                                if sz < floor:
                                    res["floors"].append(f"s{no} {nm}: {sz} pt '{r.findtext('a:t', namespaces=NS)[:30]}'")
                                runs.append((r.findtext("a:t", namespaces=NS) or "", sz, rp is not None and rp.get("b") == "1"))
                            szmax = max([r[1] for r in runs] or [floor])
                            hh += wrap_lines(runs, w * 0.97) * szmax * 1.17 / 72
                        hmax = max(hmax, hh + mt + mb)
                    tot += max(hmax, Emu(row.height).inches)
                bottom = Emu(sp.top).inches + tot
                fit_notes.append(f"s{no} {nm}: ends at y {bottom:.2f}")
                blocks.append([Emu(sp.left).inches, Emu(sp.top).inches, Emu(sp.left).inches + Emu(sp.width).inches, bottom])
                has_table = True
                if G and bottom > G["rule"][1] - 0.02:
                    res["fit"].append(f"s{no} {nm}: table ends at y {bottom:.2f}, below the area")
                continue
            tb = el.find("p:txBody", NS)
            if tb is None:
                continue
            text = "".join(t.text or "" for t in tb.iter("{%s}t" % NS["a"]))
            if not text.strip():
                continue
            ph = el.find(".//p:nvPr/p:ph", NS)
            ptype = None if ph is None else (ph.get("type") or "body")
            if ptype == "sldNum":
                continue
            # tags
            if tags.search(text):
                res["tags"].append(f"s{no} {nm}: {text[:60]!r}")
            for r_ in restricted:
                if r_["text"].lower() in text.lower() and not nm.startswith(tuple(r_.get("allowIn", []))):
                    res["tags"].append(f"s{no} {nm}: {r_['text']!r} outside {r_.get('allowIn')}")
            # floors and fit
            floor = 16 if nm in ("band", "stepcaption") else 12
            box = box_of(el, inh)
            wrap = body_attr(el, inh, "wrap", "square") != "none"
            lIns = emu_in(body_attr(el, inh, "lIns", 91440))
            rIns = emu_in(body_attr(el, inh, "rIns", 91440))
            tIns = emu_in(body_attr(el, inh, "tIns", 45720))
            bIns = emu_in(body_attr(el, inh, "bIns", 45720))
            need = 0.0
            nlines_total = 0
            inkw, em = 0.0, []  # widest paragraph; size of each paragraph's largest run
            algn = None
            for i, p in enumerate(tb.findall("a:p", NS)):
                runs = []
                for r in p.findall("a:r", NS):
                    sz, b = run_size(r, el, inh, ptype), run_bold(r, el, inh)
                    if sz < floor and not (lay == "Title Opener" and sz >= 12):
                        res["floors"].append(f"s{no} {nm}: {sz} pt '{(r.findtext('a:t', namespaces=NS) or '')[:30]}'")
                    rt = r.findtext("a:t", namespaces=NS) or ""
                    rpr = r.find("a:rPr", NS)
                    cap = (rpr.get("cap") if rpr is not None else None) or lst_val(el, inh, "/a:defRPr", "cap")
                    runs.append((rt.upper() if cap == "all" else rt, sz, b))  # an all-caps style (kickers) is measured in capitals
                if not runs:
                    continue
                ppr = p.find("a:pPr", NS)
                marL = ppr.get("marL") if ppr is not None and ppr.get("marL") else lst_val(el, inh, "", "marL")
                marL = emu_in(marL) if marL else 0.0
                lnpct = None
                ln = p.find("a:pPr/a:lnSpc/a:spcPct", NS)
                if ln is not None:
                    lnpct = int(ln.get("val")) / 100000
                if lnpct is None:
                    v = lst_val(el, inh, "/a:lnSpc/a:spcPct", "val")
                    lnpct = int(v) / 100000 if v else 1.0
                bef = 0.0
                sb = p.find("a:pPr/a:spcBef/a:spcPts", NS)
                v = sb.get("val") if sb is not None else lst_val(el, inh, "/a:spcBef/a:spcPts", "val")
                if v and i > 0:
                    bef = int(v) / 100 / 72
                szmax = max(r[1] for r in runs)
                n = wrap_lines(runs, (box[2] - lIns - rIns - marL) * 0.97) if (wrap and box) else 1
                pw = marL + sum(width_in(t, sz_, b_) for t, sz_, b_ in runs)
                if wrap and box:
                    pw = min(pw, box[2] - lIns - rIns)
                inkw, algn = max(inkw, pw), algn or (ppr.get("algn") if ppr is not None else None) or lst_val(el, inh, "", "algn")
                em.append(szmax / 72)
                nlines_total += n
                need += bef + n * szmax * 1.17 * lnpct / 72
            if ptype == "title" and lay != "Title Opener" and nlines_total > 1:
                res["fit"].append(f"s{no} title wraps to {nlines_total} lines: {text!r}")
            if box and wrap and ptype != "title" and need > box[3] - tIns - bIns + 0.03:
                res["fit"].append(f"s{no} {nm}: needs {need:.2f} in, box {box[3] - tIns - bIns:.2f} in ({nlines_total} lines): {text[:50]!r}")
            if nm == "band":
                fit_notes.append(f"s{no} band {nlines_total} lines")
            # the ink of the text: measured width placed by its alignment, the line block placed by the anchor,
            # trimmed to the glyphs (cap top to baseline-ish) so a box's empty padding never counts as overlap
            if box and em:
                x0 = box[0] + lIns
                avail = box[2] - lIns - rIns
                if algn == "ctr":
                    x0 += (avail - inkw) / 2
                elif algn == "r":
                    x0 += avail - inkw
                anchor = body_attr(el, inh, "anchor", "t")
                top = box[1] + tIns
                if anchor == "ctr":
                    top = box[1] + tIns + (box[3] - tIns - bIns - need) / 2
                elif anchor == "b":
                    top = box[1] + box[3] - bIns - need
                ink = [x0, top + 0.2 * em[0], x0 + inkw, top + need - 0.12 * em[-1]]
                texts.append((ink, text, nm))
                if not is_ph and not in_col:
                    blocks.append(ink)
        # overlap: text on text FAILs; a line through a text WARNs (gridlines, arrows and curves over labels)
        for i in range(len(texts)):
            a = texts[i][0]
            for j in range(i + 1, len(texts)):
                b = texts[j][0]
                if min(a[2], b[2]) - max(a[0], b[0]) > OVERLAP_MIN and min(a[3], b[3]) - max(a[1], b[1]) > OVERLAP_MIN:
                    res["overlap"].append(f"s{no}: {texts[i][1][:28]!r} on {texts[j][1][:28]!r}")
            core = [a[0] + 0.01, a[1] + 0.01, a[2] - 0.01, a[3] - 0.01]
            for x1, y1, x2, y2, ln_nm in segs:
                if core[2] > core[0] and core[3] > core[1] and seg_hits_rect(x1, y1, x2, y2, core):
                    warns["lines"].append(f"s{no}: line {ln_nm} crosses {texts[i][1][:28]!r}")
            # a label that runs out of the shape it sits in (partly inside, partly outside by > OVERLAP_MIN)
            for r, bnm in boxes:
                ix = min(a[2], r[2]) - max(a[0], r[0])
                iy = min(a[3], r[3]) - max(a[1], r[1])
                out_ = max(r[0] - a[0], a[2] - r[2], r[1] - a[1], a[3] - r[3])
                if ix > OVERLAP_MIN and iy > OVERLAP_MIN and out_ > OVERLAP_MIN and (ix * iy) > 0.3 * (a[2] - a[0]) * (a[3] - a[1]):
                    warns["edges"].append(f"s{no}: {texts[i][1][:28]!r} runs over the edge of {bnm}")
        # white space: the share of the content (or frame) area that no element reaches
        if lay in areas:
            ef = empty_fraction(areas[lay], blocks, [s_[:4] for s_ in segs])
            kind = "frame" if lay == "Header w/storyboard" else "table" if has_table else "content"
            ws_notes.append(f"s{no} {ef:.1%}")
            if ef > WS_WARN[kind]:
                warns["whitespace"].append(f"s{no}: {ef:.0%} of the {kind} area is empty (warn above {WS_WARN[kind]:.0%})")

    want = sum(1 for m in meta["slides"] if m.get("morph")) if meta else morph_n
    if morph_n != want:
        res["morph"].append(f"{morph_n} Morph transitions, want {want}")
    # speaker notes on every slide
    for no, slide in enumerate(prs.slides, start=1):
        t = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        if not t:
            res["notes"].append(f"s{no}: no speaker notes")
    # deck structure (from the sidecar's slide kinds)
    if meta and all("kind" in m for m in meta["slides"]):
        ms = meta["slides"]
        kinds = [m["kind"] for m in ms]
        if kinds[0] != "title":
            res["structure"].append(f"s1 is {kinds[0]}, want the title slide")
        if len(kinds) < 2 or kinds[1] != "outline":
            res["structure"].append("s2 is not the outline")
        # the Summary part right after the title: the outline opens it (as its divider), the summary slide follows
        elif not ms[1]["section"].startswith("Summary"):
            res["structure"].append(f"s2 (the outline) opens part '{ms[1]['section']}', want the Summary part")
        elif len(ms) < 3 or ms[2]["section"] != ms[1]["section"] or ms[2]["kind"] not in ("content", "frame-first"):
            res["structure"].append("s3 is not the summary slide (want title, outline, summary)")
        else:
            t3 = prs.slides[2].shapes.title.text.strip() if prs.slides[2].shapes.title is not None else ""
            if not t3.startswith("Summary"):
                res["structure"].append(f"s3 title {t3[:40]!r} does not start with 'Summary' (want 'Summary: <core message>')")
        closing = [i for i, k in enumerate(kinds) if k == "closing"]
        if len(closing) != 1:
            res["structure"].append(f"{len(closing)} closing slides, want 1")
        else:
            c = closing[0]
            for i, m in enumerate(ms):
                if m.get("appendix") and i < c:
                    res["structure"].append(f"s{i + 1}: appendix slide before the closing slide")
                if i > c and not m.get("appendix"):
                    res["structure"].append(f"s{i + 1}: main slide after the closing slide (mark it appendix or move it)")
        first_section = ms[0]["section"]
        for i, m in enumerate(ms):
            new_sec = i == 0 or ms[i - 1]["section"] != m["section"]
            if new_sec and m["section"] != first_section and not m.get("appendix") and m["kind"] not in ("divider", "outline"):
                res["structure"].append(f"s{i + 1}: part '{m['section']}' opens without a divider")
    # storyboards: consecutive frame slides with one title; the unkeyed !!u<n> elements must match
    prev = None
    for no in sorted(story_names):
        title, names = story_names[no]
        if prev and prev[0] == no - 1 and prev[1] == title:
            a = {n for n in prev[2] if re.fullmatch(r"!!u\d+", n)}
            b = {n for n in names if re.fullmatch(r"!!u\d+", n)}
            if a != b:
                res["morph"].append(f"s{no}: unkeyed !! set differs from s{no - 1}")
            if len(names) != len(set(names)):
                res["morph"].append(f"s{no}: duplicate !! names")
        prev = (no, title, names)

    ok = True
    for k, v in res.items():
        print(f"{'PASS' if not v else 'FAIL'} {k}" + ("" if not v else ": " + str(len(v))))
        for x in v:
            print("   ", x)
        ok &= not v
    for k, v in warns.items():
        print(f"{'PASS' if not v else 'WARN'} {k}" + ("" if not v else ": " + str(len(v))))
        for x in v:
            print("   ", x)
    print("info:", "; ".join(fit_notes))
    if ws_notes:
        print("info: empty share of the content / frame area:", " ".join(ws_notes))
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
